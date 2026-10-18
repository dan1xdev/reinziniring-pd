"""Полный прогон и одна итоговая строка на ИНН."""
import argparse
from collections import defaultdict
from datetime import datetime, timezone
import getpass
import json
import os
from pathlib import Path
import re
from openpyxl import Workbook
from .dadata import DaDataClient
from .export import add_table, RECORD_COLUMNS
from .input import read_records
from .normalization import normalize
from .pipeline import process
from .resolution import RegistrySearch, SearchCache, resolve


def aggregate(result):
    records = {r['record_id']: r for r in result['records']}
    buckets = defaultdict(list)
    review = []
    candidates_by_group = defaultdict(list)
    for candidate in result['resolution'].get('candidates', []):
        candidates_by_group[candidate['group_id']].append(candidate)
    for resolution in result['resolution']['groups']:
        selected = resolution['selected']
        members = [records[i] for i in resolution['record_ids']]
        if selected:
            buckets[selected['candidate_inn']].append((resolution, members))
        else:
            options = sorted(candidates_by_group[resolution['group_id']],
                             key=lambda c: (c['eligible'], c['score']), reverse=True)
            top = {}
            for option in options:
                top.setdefault(option['candidate_inn'], option)
            top = list(top.values())[:3]
            reasons = sorted({reason for option in top for reason in option['reasons']})
            review.append({**resolution, 'region': members[0]['region'], 'city': members[0]['city'],
                'students': '\n'.join(dict.fromkeys(r.get('student_name', '') for r in members if r.get('student_name', ''))),
                'files': '\n'.join(dict.fromkeys(r['personal_file_number'] for r in members if r['personal_file_number'])),
                'reason': '\n'.join(reasons) or ('Несколько подходящих ИНН' if options else 'Кандидаты не найдены'),
                'top_candidates': '\n\n'.join(f"{c['candidate_inn']} | {c['candidate_name']} | {c['address']} | {c['source']}" for c in top),
                'record_ids': '\n'.join(resolution['record_ids'])})
        for r in members:
            r['resolved_inn'] = selected['candidate_inn'] if selected else ''
            r['resolution_status'] = resolution['status']
            r['resolution_source'] = selected['source'] if selected else ''
    schools, students = [], []
    for inn, groups in buckets.items():
        members = [r for _, group in groups for r in group]
        selected = [g['selected'] for g, _ in groups]
        preferred = [c for c in selected if c['source'] == 'DaData'] or selected
        candidate = preferred[0]
        addresses = {c['address'] for c in preferred if c['address']}
        address = next(iter(addresses)) if len(addresses) == 1 else ''
        address = re.sub(r'^\s*\d{6}\s*,?\s*', '', address)
        address = re.sub(r'^Россия\s*,\s*', '', address, flags=re.I)
        pairs = {}
        for r in members:
            name, file = r.get('student_name', ''), r['personal_file_number']
            if name or file:
                # Неполные пары нельзя автоматически считать одним учеником.
                key = (normalize(name), file) if name and file else (r['record_id'],)
                pairs.setdefault(key, {'name': name, 'file': file, 'record_ids': []})['record_ids'].append(r['record_id'])
        people = sorted(pairs.values(), key=lambda p: (p['name'].casefold(), p['file']))
        names_per_file = defaultdict(set)
        for p in people:
            if p['file']:
                names_per_file[p['file']].add(normalize(p['name']))
            students.append({'school_name': candidate['candidate_name'], 'inn': inn,
                'address': address, 'student': p['name'], 'file': p['file'], 'record_ids': '\n'.join(p['record_ids'])})
        notes = []
        if any(len(names) > 1 for names in names_per_file.values()):
            notes.append('Один номер дела связан с разными ФИО — проверить')
        if len(addresses) != 1:
            notes.append('Адрес отсутствует или есть противоречие — проверить')
        if candidate['source'] != 'DaData':
            notes.append('Адрес реестра аккредитации; актуальность юридического адреса не подтверждена')
        missing = sum(not r.get('student_name', '') or not r['personal_file_number'] for r in members)
        if missing:
            notes.append(f'В {missing} исходных строках нет ФИО и/или номера дела')
        schools.append({'school_name': candidate['candidate_name'], 'address': address,
            'students': '\n'.join(p['name'] or 'ФИО отсутствует' for p in people),
            'files': '\n'.join(p['file'] or 'Номер отсутствует' for p in people),
            'student_count': len(people), 'inn': inn, 'ogrn': candidate['candidate_ogrn'],
            'postal_code': candidate['postal_code'], 'region': candidate['region'],
            'source': candidate['source'], 'status': 'Сопоставлено автоматически',
            'notes': '\n'.join(notes), 'aliases': '\n'.join(sorted({r['name'] for r in members})),
            'row_count': len(members), 'record_ids': '\n'.join(r['record_id'] for r in members)})
    schools.sort(key=lambda r: (not bool(r['student_count']), r['school_name'].casefold()))
    result.update(schools=schools, students=students, review=review)
    groups = result['resolution']['groups']
    result['full_summary'] = {
        'Входных строк': len(result['records']), 'Отобрано строк': result['summary']['selected_rows'],
        'Сопоставлено исходных строк': sum(r['row_count'] for r in schools),
        'Исходных строк на проверку': sum(len(g['record_ids']) for g in groups if not g['selected']),
        'Учреждений без повторов по ИНН': len(schools),
        'Учреждений с адресом': sum(bool(r['address']) for r in schools),
        'Учреждений с данными учеников': sum(bool(r['student_count']) for r in schools),
        'Пар ФИО и номер дела после объединения': len(students), 'Групп на проверку': len(review),
        'Групп сопоставлено через DaData': sum(bool(g['selected']) and g['selected']['source'] == 'DaData' for g in groups),
        'Групп сопоставлено через реестр': sum(bool(g['selected']) and g['selected']['source'] != 'DaData' for g in groups),
        'Групп проверено по реестру': sum(g['registry_used'] for g in groups),
        'Запросов DaData в этом запуске': result['resolution']['actual_requests'],
        'Успешных уникальных запросов с учётом кэша': len({r['query'] for r in result['resolution']['requests'] if r['status'] == 'OK'}),
        'Ошибка или пропуск запроса': sum(r['status'] != 'OK' for r in result['resolution']['requests']),
        'Строк без данных ученика': sum(not r.get('student_name', '') or not r['personal_file_number'] for r in result['records']),
    }
    return result


def save(result, output):
    output.mkdir(parents=True, exist_ok=True)
    (output / 'result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2))
    book = Workbook()
    book.remove(book.active)
    main = add_table(book, 'Учреждения', result['schools'], [
        ('school_name', 'Школа / учреждение'), ('address', 'Адрес организации'),
        ('students', 'ФИО учеников'), ('files', 'Номера личных дел'), ('student_count', 'Учеников / дел'),
        ('postal_code', 'Индекс'), ('inn', 'ИНН'), ('source', 'Источник адреса'),
        ('notes', 'Примечания'), ('region', 'Регион'), ('ogrn', 'ОГРН'), ('status', 'Статус сопоставления'),
        ('row_count', 'Исходных строк')])
    for col, width in {'A':65,'B':55,'C':35,'D':23,'E':15,'F':12,'G':18,'H':28,'I':65}.items():
        main.column_dimensions[col].width = width
    add_table(book, 'Ученики и дела', result['students'], [('school_name','Школа / учреждение'),
        ('address','Адрес'),('student','ФИО ученика'),('file','Номер личного дела'),('inn','ИНН'),('record_ids','Исходные записи')])
    review_sheet = add_table(book, 'Проверить', result['review'], [('source_name','Исходное название'),('region','Регион'),
        ('city','Город'),('status','Статус'),('students','ФИО'),('files','Номер дела'),('reason','Что уточнить'),
        ('top_candidates','Первые три кандидата'),('candidate_count','Кандидатов'),('registry_used','Проверен реестр'),('registry_truncated','Выдача ограничена'),
        ('group_id','Группа'),('record_ids','Исходные записи')])
    review_sheet.column_dimensions['H'].width = 110
    for row in range(2, review_sheet.max_row + 1):
        review_sheet.row_dimensions[row].height = 240 if result['review'][row-2]['top_candidates'] else 65
    summary = [{'metric':k,'value':v} for k,v in result['full_summary'].items()]
    summary += [{'metric':'Единица группировки','value':'Одно юридическое лицо (ИНН), независимо от написания. Корпуса не разделены.'},
        {'metric':'Адрес','value':'DaData: юридический адрес. Реестр: адрес организации из реестра аккредитации, требует сверки.'},
        {'metric':'ФИО и номера дел','value':'Строки в соседних ячейках совпадают по порядку. Полные пары удалены как дубли только внутри одного ИНН.'},
        {'metric':'Дата обработки UTC','value':result['created_at_utc']},
        {'metric':'Дата снимка реестра','value':result['resolution']['registry'].get('snapshot_date', '')},
        {'metric':'DaData','value':'https://dadata.ru/api/suggest/party/'},
        {'metric':'Резервный реестр','value':'https://isga.obrnadzor.gov.ru/accredreestr/opendata/'}]
    add_table(book,'Сводка',summary,[('metric','Показатель'),('value','Значение')])
    add_table(book,'Исходные записи',result['records'],[('resolved_inn','Сопоставленный ИНН'),
        ('resolution_status','Результат сопоставления'),('resolution_source','Источник сопоставления'),*RECORD_COLUMNS])
    add_table(book,'Запросы DaData',result['resolution']['requests'],[('group_id','Группа'),('query','Запрос'),
        ('status','Статус'),('candidate_count','Кандидатов'),('cached','Из кэша'),('timestamp_utc','Время UTC')])
    # В книге достаточно первых пяти различных ИНН на группу; полный аудит есть в JSON.
    candidate_groups = defaultdict(list)
    for c in result['resolution']['candidates']:
        candidate_groups[c['group_id']].append(c)
    shortlist = []
    for group in candidate_groups.values():
        group.sort(key=lambda c: (c['eligible'], c['score'], c['legal_status'] in {'ACTIVE','Действующее'}), reverse=True)
        distinct = {}
        for c in group:
            distinct.setdefault(c['candidate_inn'], c)
        shortlist.extend(list(distinct.values())[:5])
    add_table(book,'Кандидаты',shortlist,[('group_id','Группа'),('source','Источник'),
        ('candidate_name','Название'),('candidate_inn','ИНН'),('address','Адрес'),('legal_status','Статус источника'),
        ('eligible','Прошёл правила'),('score','Совпадение слов, %'),('reasons','Причины отказа'),
        ('source_url','Источник в интернете'),('certificate_id','ID записи реестра')])
    add_table(book,'Источники',result['sources'],[('file','Файл'),('rows','Строк'),('sha256','SHA-256'),('source_id','Код источника')])
    book.save(output / 'schools.xlsx')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('inputs',nargs='+',type=Path)
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--cache',type=Path,required=True)
    parser.add_argument('--registry',type=Path,default=Path(__file__).resolve().parents[2]/'registry_cache')
    parser.add_argument('--online',action='store_true')
    parser.add_argument('--ask-token',action='store_true')
    args = parser.parse_args()
    if args.output.exists() and any(args.output.iterdir()):
        parser.error('Папка результата должна быть новой или пустой')
    if args.ask_token and not args.online:
        parser.error('--ask-token требует --online')
    config = json.loads((Path(__file__).resolve().parents[1]/'config.json').read_text())
    records,sources = read_records(args.inputs)
    result = process(records,config)
    result.update(sources=sources,created_at_utc=datetime.now(timezone.utc).isoformat())
    token = getpass.getpass('API-токен DaData: ') if args.ask_token else os.environ.get('DADATA_API_KEY','')
    if args.online and not token:
        parser.error('Нужен DADATA_API_KEY или --ask-token')
    api = SearchCache(args.cache,DaDataClient(token) if args.online else None)
    registry = RegistrySearch(args.registry)
    resolve(result,api,registry)
    aggregate(result)
    save(result,args.output)
    print(json.dumps(result['full_summary'],ensure_ascii=False,indent=2))
    return 2 if result['full_summary']['Ошибка или пропуск запроса'] else 0


if __name__ == '__main__':
    raise SystemExit(main())
