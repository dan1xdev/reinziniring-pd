"""Экспорт результата Python-прототипа в JSON и отдельную книгу Excel."""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .roster import attach_roster


RECORD_COLUMNS = [
    ("record_id", "ID записи"), ("case_id", "CASE"), ("selection", "Отбор"),
    ("selection_reason", "Причина отбора"), ("name", "Исходное название"),
    ("normalized_name", "Нормализованное название"), ("search_variant", "Поисковый вариант"),
    ("education_kind", "Вид образования"), ("education_source", "Источник вида образования"),
    ("region", "Регион"), ("city", "Город"), ("region_evidence", "Основания региона"),
    ("city_evidence", "Основания города"), ("region_hypothesis", "Гипотеза региона"),
    ("flags", "Признаки риска"), ("group_id", "Предварительная группа"), ("status", "Статус"),
    ("confirmed_inn", "Подтвержденный ИНН"), ("official_name", "Официальное название"),
    ("region_raw", "Исходный регион"), ("location_raw", "Исходное местонахождение"),
    ("year", "Год завершения"), ("source_file", "Исходный файл"),
    ("source_sheet", "Исходный лист"), ("source_row", "Строка в книге"), ("original_row", "Строка исходного реестра"),
    ("personal_file_number", "Номер личного дела"),
    ("student_name", "ФИО ученика"), ("school_address_raw", "Адрес ОУ из источника"),
    ("postal_code_raw", "Индекс из источника"),
]


def excel_value(value):
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return value


def add_table(workbook, name, rows, columns):
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.append([label for _, label in columns])
    for row in rows:
        sheet.append(["\n".join(row.get(key, [])) if key == "personal_file_numbers"
                      else excel_value(row.get(key, "")) for key, _ in columns])
    for cells in sheet:
        for cell in cells:
            # Все строки — литералы, включая текст с начальным знаком равенства.
            if isinstance(cell.value, str):
                cell.data_type = "s"
                cell.number_format = "@"
            cell.font = Font(name="Arial", size=10, color="202936")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="273B53")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, (key, _) in enumerate(columns, 1):
        long = any(word in key for word in ["name", "query", "reason", "evidence", "flags", "address", "variant"])
        sheet.column_dimensions[get_column_letter(index)].width = 60 if long else 27
    widths = {
        "Сводка": {"A": 42, "B": 65},
        "Источники": {"A": 42, "D": 68, "E": 70},
        "Запросы DaData": {"C": 37, "E": 35},
        "Группы": {"A": 65, "B": 25, "C": 18, "D": 20, "G": 38, "H": 27, "L": 65},
        "Школы и ученики": {"A": 60, "B": 60, "C": 35, "D": 32, "E": 23, "F": 14,
                              "G": 40, "H": 32, "I": 25, "J": 60, "K": 25, "L": 45},
    }
    for column, width in widths.get(name, {}).items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 32
    for row_index, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(columns)), 2):
        max_lines = max((sum(max(1, (len(part) + int(sheet.column_dimensions[cell.column_letter].width) - 1)
                                    // int(sheet.column_dimensions[cell.column_letter].width))
                            for part in str(cell.value or "").splitlines()) for cell in cells), default=1)
        sheet.row_dimensions[row_index].height = max(32, min(180, 15 * max_lines + 8))
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def write_result(result, destination):
    attach_roster(result)
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    json_path = destination / "result.json"
    json_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    workbook = Workbook()
    workbook.remove(workbook.active)
    summary = result["summary"]
    labels = {"input_rows": "Входных строк", "selected_rows": "Отобрано", "excluded_rows": "Исключено",
              "selection_review_rows": "Требуют уточнения отбора", "preliminary_groups": "Предварительных групп",
              "possible_matches": "Пар возможных совпадений", "confirmed_organizations": "Подтвержденных организаций"}
    labels.update({"selected_rows_with_personal_file": "Отобрано строк с номером личного дела",
                   "selected_rows_without_personal_file": "Отобрано строк без номера личного дела",
                   "distinct_personal_file_numbers": "Различных номеров личных дел"})
    labels.update({"selected_rows_with_student_name": "Отобрано строк с ФИО",
                   "selected_rows_without_student_name": "Отобрано строк без ФИО",
                   "groups_with_address_available": "Групп с адресом для проверки",
                   "groups_without_address": "Групп без определенного адреса"})
    overview = [{"metric": label, "value": summary[key]} for key, label in labels.items()]
    overview.extend([{"metric": "Режим DaData", "value": result["dadata"]["mode"]},
                     {"metric": "Реальных запросов DaData", "value": result["dadata"]["actual_requests"]},
                     {"metric": "Назначение результата", "value": "Предварительная обработка; ИНН не подтверждены"}])
    add_table(workbook, "Сводка", overview, [("metric", "Показатель"), ("value", "Значение")])
    add_table(workbook, "Записи", result["records"], RECORD_COLUMNS)
    groups_for_display = sorted(result["groups"], key=lambda g: (not bool(g["personal_file_numbers"]),
                                                               g["normalized_name"], g["group_id"]))
    add_table(workbook, "Группы", groups_for_display, [("name", "Образовательное учреждение"),
              ("personal_file_numbers", "Номера личных дел"), ("personal_file_count", "Различных номеров"),
              ("rows_without_personal_file", "Строк без номера"), ("region", "Регион"), ("city", "Город"),
              ("personal_file_status", "Полнота номеров"), ("status", "Статус группировки"),
              ("group_id", "Группа"), ("row_count", "Входных строк"),
              ("basis", "Основание"), ("record_ids", "ID записей")])
    workbook.move_sheet("Группы", offset=-2)
    add_table(workbook, "Школы и ученики", result["school_students"], [
        ("school_name", "Школа / учреждение"), ("school_address", "Адрес организации"),
        ("address_status", "Статус адреса"), ("student_name", "ФИО ученика"),
        ("personal_file_number", "Номер личного дела"), ("postal_code", "Индекс"),
        ("address_source", "Источник адреса"), ("student_status", "Полнота данных ученика"),
        ("candidate_inn", "ИНН кандидата"), ("candidate_name", "Название кандидата DaData"),
        ("group_id", "Предварительная группа"), ("source_record_ids", "ID исходных записей")])
    workbook.move_sheet("Школы и ученики", offset=-3)
    add_table(workbook, "Возможные совпадения", result["possible_matches"], [
        ("left_group", "Группа 1"), ("right_group", "Группа 2"), ("left_name", "Название 1"),
        ("right_name", "Название 2"), ("text_similarity", "Сходство текста, %"),
        ("region", "Регион"), ("reason", "Основание"), ("status", "Статус")])
    add_table(workbook, "Запросы DaData", result["dadata"]["requests"], [
        ("group_id", "Группа"), ("query", "Запрос"), ("status", "Результат запроса"),
        ("candidate_count", "Кандидатов"), ("timestamp_utc", "Время UTC")])
    add_table(workbook, "Кандидаты DaData", result["dadata"]["candidates"], [
        ("group_id", "Группа"), ("rank", "Позиция в ответе"), ("candidate_name", "Название кандидата"),
        ("candidate_inn", "ИНН кандидата"), ("candidate_kpp", "КПП"), ("candidate_ogrn", "ОГРН"),
        ("legal_status", "Статус юрлица"), ("branch_type", "Тип филиала"), ("address", "Адрес кандидата"),
        ("postal_code", "Индекс"), ("flags", "Признаки риска"), ("decision", "Решение")])
    add_table(workbook, "Источники", result["sources"], [("file", "Файл"), ("source_id", "Код источника"),
        ("rows", "Прочитано строк"), ("sha256", "SHA-256"), ("skipped_sheets", "Пропущенные листы")])
    path = destination / "result.xlsx"
    workbook.save(path)
    return path
