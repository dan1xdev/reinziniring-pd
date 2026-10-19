"""Чтение двух предоставленных форматов без использования аналитической разметки."""

import hashlib
from pathlib import Path

from openpyxl import load_workbook

from .normalization import normalize, text


ALIASES = {
    "name": ["Образовательное учреждение", "Название ОУ в исходном Excel"],
    "education": ["Вид образования"],
    "document": ["Тип документа об образовании"],
    "region_raw": ["Регион (как в файле)", "Регион", "Регион ОУ"],
    "location_raw": ["Местонахождение ОУ", "Местонахождение"],
    "city_raw": ["Населённый пункт", "Населенный пункт", "Город"],
    "year": ["Год завершения", "Год окончания"],
    "case_id": ["ID"],
    "original_row": ["Строка в Excel"],
    "inn_raw": ["ИНН", "ИНН ОУ"],
    "personal_file_number": ["Личное дело.Номер", "Номер личного дела", "Личное дело"],
    "student_name": ["ФИО", "ФИО ученика", "ФИО поступающего"],
    "school_address_raw": ["Адрес", "Адрес ОУ", "Адрес школы"],
    "postal_code_raw": ["Индекс", "Почтовый индекс"],
}


def personal_file_number(cell):
    # Сохраняем ведущие нули, в том числе заданные простым числовым форматом Excel.
    value = cell.value
    if isinstance(value, (int, float)) and not isinstance(value, bool) and value == int(value):
        number_format = cell.number_format
        if number_format and set(number_format) == {"0"}:
            return str(int(value)).zfill(len(number_format))
    return text(value)


def read_records(paths, should_cancel=None):
    records, sources = [], []
    seen_paths = set()
    for filename in paths:
        if should_cancel:
            should_cancel()
        path = Path(filename).expanduser().resolve()
        if path in seen_paths:
            continue
        seen_paths.add(path)
        if path.suffix.lower() != ".xlsx":
            raise ValueError(f"Поддерживается .xlsx: {path.name}")
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        source_id = hashlib.sha256(str(path).encode()).hexdigest()[:8]
        workbook = load_workbook(path, read_only=True, data_only=False)
        count = 0
        skipped = []
        try:
            for sheet in workbook:
                # Аналитические выводы и «Регион для поиска» не являются входными данными.
                if "100 записей" in workbook.sheetnames and sheet.title != "100 записей":
                    skipped.append(sheet.title)
                    continue
                found = read_sheet(sheet, path.name, source_id, should_cancel)
                if found is None:
                    skipped.append(sheet.title)
                    continue
                records.extend(found)
                count += len(found)
        finally:
            workbook.close()
        if not count:
            raise ValueError(f"{path.name}: нет строк с распознаваемой структурой")
        sources.append({"file": path.name, "source_id": source_id, "sha256": digest,
                        "rows": count, "skipped_sheets": skipped})
    return records, sources


def read_sheet(sheet, filename, source_id, should_cancel=None):
    mapping = None
    records = []
    for row_number, cells in enumerate(sheet.iter_rows(), 1):
        if should_cancel and row_number % 100 == 1:
            should_cancel()
        values = [cell.value for cell in cells]
        if mapping is None:
            if row_number > 30:
                return None
            headers = [normalize(value) for value in values]
            candidate = {}
            for field, aliases in ALIASES.items():
                indices = [i for i, header in enumerate(headers)
                           if header in {normalize(alias) for alias in aliases}]
                if len(indices) > 1:
                    raise ValueError(f"{filename}/{sheet.title}: дубли колонок {field}")
                if indices:
                    candidate[field] = indices[0]
            if "name" in candidate and ("education" in candidate or "document" in candidate):
                mapping = candidate
            continue
        if all(value is None or text(value) == "" for value in values):
            continue
        record = {field: text(values[index]) for field, index in mapping.items()}
        record["name"] = "" if values[mapping["name"]] is None else str(values[mapping["name"]])
        record["personal_file_number"] = (
            personal_file_number(cells[mapping["personal_file_number"]])
            if "personal_file_number" in mapping else ""
        )
        if "postal_code_raw" in mapping:
            record["postal_code_raw"] = personal_file_number(cells[mapping["postal_code_raw"]])
        formula_fields = [field for field, index in mapping.items() if cells[index].data_type == "f"]
        for field in formula_fields:
            record[field] = ""
        record.update({"record_id": f"{source_id}:{sheet.title}:{row_number}",
                       "source_file": filename, "source_sheet": sheet.title,
                       "source_row": row_number, "formula_fields": formula_fields})
        records.append(record)
    return records if mapping is not None else None
