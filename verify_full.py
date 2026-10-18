"""Контроль сохранности присланных 140 строк и 10 пар ученик—дело."""
import json
from pathlib import Path
import sys
from openpyxl import load_workbook

root = Path(sys.argv[1])
r = json.loads((root/'result.json').read_text())
assert len(r['records']) == 140
selected = [row for row in r['records'] if row['selection'] == 'ОТОБРАНО']
assert len(selected) == 140
assert len({s['inn'] for s in r['schools']}) == len(r['schools'])
assert all(s['inn'] and s['address'] for s in r['schools'])
matched_ids = [i for s in r['schools'] for i in s['record_ids'].splitlines()]
review_ids = [i for s in r['review'] for i in s['record_ids'].splitlines()]
assert len(matched_ids) + len(review_ids) == 140
assert len(set(matched_ids + review_ids)) == 140
assert set(matched_ids + review_ids) == {s['record_id'] for s in selected}
student_rows = [row for row in selected if row.get('student_name')]
assert len(student_rows) == 40
assert all(row['resolved_inn'] for row in student_rows)
assert len(r['students']) == 10
assert len({s['inn'] for s in r['students']}) == 10
assert {s['file'] for s in r['students']} == {f'T-{i}' for i in range(1001,1011)}
assert sorted(i for s in r['students'] for i in s['record_ids'].splitlines()) == sorted(s['record_id'] for s in student_rows)
for s in r['students']:
    members = [row for row in student_rows if row['record_id'] in s['record_ids'].splitlines()]
    assert len(members) == 4
    assert {row['personal_file_number'] for row in members} == {s['file']}
    assert {row['resolved_inn'] for row in members} == {s['inn']}
    assert {row['student_name'] for row in members} == {s['student']}
assert all(q['status'] == 'OK' for q in r['resolution']['requests'])
assert {q['group_id'] for q in r['resolution']['requests']} == {g['group_id'] for g in r['groups']}
assert all(c['group_id'] for c in r['resolution']['candidates'])
queries='\n'.join(q['query'] for q in r['resolution']['requests'])
assert all(s['student'] not in queries and s['file'] not in queries for s in r['students'])
for g in r['resolution']['groups']:
    if not g['selected']:
        assert g['registry_used']
    else:
        assert any(c['group_id']==g['group_id'] and c['candidate_inn']==g['selected']['candidate_inn'] and c['eligible'] for c in r['resolution']['candidates'])
book=load_workbook(root/'schools.xlsx')
assert book['Учреждения'].max_row == len(r['schools'])+1
assert book['Ученики и дела'].max_row == 11
for sheet in book:
    assert sheet.freeze_panes and sheet.auto_filter.ref
    assert all(cell.data_type not in {'f','e'} for row in sheet for cell in row)
for row in book['Учреждения'].iter_rows(min_row=2):
    assert row[6].data_type == 's'
    assert row[6].number_format == '@'
print('OK: все 140 строк учтены; 40 повторных записей сведены к 10 связанным парам; ИНН и адреса заполнены; ошибок API нет.')
