"""Проверка результата именно на пяти приложенных тестовых книгах."""

import json
from pathlib import Path
import sys

from openpyxl import load_workbook


def verify(folder):
    folder = Path(folder)
    data = json.loads((folder / "result.json").read_text(encoding="utf-8"))
    rows = data["records"]
    assert len(rows) == 140, "Ожидались 40 вариантов и 100 аналитических случаев"
    assert len({r["record_id"] for r in rows}) == 140, "Повтор ID записи"
    assert sum(source["rows"] for source in data["sources"]) == 140
    assert sum(data["summary"][key] for key in ["selected_rows", "excluded_rows", "selection_review_rows"]) == 140
    by_file = {}
    for record in rows:
        by_file.setdefault(record["source_file"], []).append(record)
    ordinary = by_file["01_обычное_написание.xlsx"]
    caps = by_file["02_названия_капсом.xlsx"]
    assert len(ordinary) == len(caps) == 10
    for a, b in zip(ordinary, caps):
        assert a["normalized_name"] == b["normalized_name"]
        assert a["group_id"] == b["group_id"]
    cases = {record["case_id"]: record for record in rows if record.get("case_id")}
    assert set(cases) == {f"CASE-{i:03d}" for i in range(1, 101)}
    for i in range(1, 13, 2):
        assert cases[f"CASE-{i:03d}"]["group_id"] != cases[f"CASE-{i + 1:03d}"]["group_id"]
    for i in range(31, 45):
        assert "ПОДОЗРЕНИЕ_НА_ОБРЕЗКУ" in cases[f"CASE-{i:03d}"]["flags"]
    for i in range(67, 77):
        assert "СТАРЫЙ_ВЫПУСК" in cases[f"CASE-{i:03d}"]["flags"]
    for i in range(77, 85):
        record = cases[f"CASE-{i:03d}"]
        assert record["selection"] == "ОТОБРАНО"
        assert "ВОЗМОЖНОЕ_ПОДРАЗДЕЛЕНИЕ_ВУЗА" in record["flags"]
    members = [identifier for group in data["groups"] for identifier in group["record_ids"]]
    selected = [r["record_id"] for r in rows if r["selection"] == "ОТОБРАНО"]
    assert sorted(members) == sorted(selected), "Потеря или повтор записей в группах"
    assert all(not r["confirmed_inn"] and not r["official_name"] for r in rows)
    assert data["summary"]["selected_rows_with_personal_file"] == 40
    assert data["summary"]["selected_rows_without_personal_file"] == 100
    assert data["summary"]["distinct_personal_file_numbers"] == 10
    assert data["summary"]["selected_rows_with_student_name"] == 40
    assert data["summary"]["selected_rows_without_student_name"] == 100
    by_id = {r["record_id"]: r for r in rows}
    for group in data["groups"]:
        expected = sorted({by_id[i]["personal_file_number"] for i in group["record_ids"]
                           if by_id[i]["personal_file_number"]})
        assert group["personal_file_numbers"] == expected
        assert group["personal_file_count"] == len(expected)
    roster_sources = []
    for item in data["school_students"]:
        for identifier in item["source_record_ids"]:
            original = by_id[identifier]
            assert item["student_name"] == original.get("student_name", "")
            assert item["personal_file_number"] == original["personal_file_number"]
            assert item["group_id"] == original["group_id"]
            roster_sources.append(identifier)
        if item["school_address"]:
            assert "ПРОВЕРИТЬ" in item["address_status"]
    assert sorted(roster_sources) == sorted(selected)
    workbook = load_workbook(folder / "result.xlsx", read_only=True, data_only=False)
    assert workbook["Записи"].max_row == 141
    assert workbook.sheetnames[0] == "Школы и ученики"
    assert workbook["Школы и ученики"].max_row == len(data["school_students"]) + 1
    assert workbook["Группы"].max_row == len(data["groups"]) + 1
    assert workbook["Кандидаты DaData"].max_row == max(1, len(data["dadata"]["candidates"]) + 1)
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()
    assert all("ФИО" not in query["query"] for query in data["dadata"]["requests"])
    print("Проверка 140 входных записей, групп, рисков и экспорта: OK")
    print(f"DaData: {data['dadata']['actual_requests']} запросов; {len(data['dadata']['candidates'])} кандидатов")


if __name__ == "__main__":
    verify(sys.argv[1] if len(sys.argv) > 1 else "results")
