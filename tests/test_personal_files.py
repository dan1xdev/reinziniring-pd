"""Связь учреждения с личными делами и отсутствие этих номеров в запросах API."""

import json
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from ou_processor.dadata import run_sample
from ou_processor.export import write_result
from ou_processor.input import read_records
from ou_processor.pipeline import process


CONFIG = json.loads((Path(__file__).resolve().parents[1] / "config.json").read_text(encoding="utf-8"))


def record(index, number, **values):
    return {"record_id": str(index), "name": "ГБОУ города Москвы Школа № 1",
            "education": "Среднее общее", "personal_file_number": number,
            "formula_fields": [], **values}


class PersonalFileTests(unittest.TestCase):
    def test_aggregate_unique_numbers_and_report_missing(self):
        result = process([record(1, "0012"), record(2, "0012"), record(3, "T-9"), record(4, "")], CONFIG)
        group = result["groups"][0]
        self.assertEqual(group["personal_file_numbers"], ["0012", "T-9"])
        self.assertEqual(group["personal_file_count"], 2)
        self.assertEqual(group["row_count"], 4)
        self.assertEqual(group["rows_without_personal_file"], 1)
        self.assertEqual(group["personal_file_status"], "НОМЕРА_ЕСТЬ_ЧАСТИЧНО")

    def test_same_number_does_not_merge_schools_or_include_excluded_rows(self):
        result = process([record(1, "T-1"), record(2, "T-1", name="ГБОУ города Москвы Школа № 2"),
                          record(3, "T-2", education="Высшее")], CONFIG)
        self.assertEqual(len(result["groups"]), 2)
        self.assertTrue(all(g["personal_file_numbers"] == ["T-1"] for g in result["groups"]))
        self.assertEqual(result["summary"]["distinct_personal_file_numbers"], 1)

    def test_missing_number_not_replaced_by_case_id(self):
        result = process([record(1, "", case_id="CASE-001", original_row="2056")], CONFIG)
        self.assertEqual(result["groups"][0]["personal_file_numbers"], [])
        self.assertEqual(result["groups"][0]["personal_file_status"], "НЕТ_НОМЕРОВ_В_ИСТОЧНИКЕ")

    def test_read_and_export_leading_zeros(self):
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Образовательное учреждение", "Вид образования", "Личное дело.Номер"])
            sheet.append(["ГБОУ города Москвы Школа № 1", "Среднее общее", 12])
            sheet["C2"].number_format = "00000"
            sheet.append(["ГБОУ города Москвы Школа № 1", "Среднее общее", "00013"])
            workbook.save(source)
            rows, sources = read_records([source])
            self.assertEqual([r["personal_file_number"] for r in rows], ["00012", "00013"])
            result = process(rows, CONFIG)
            result["sources"] = sources
            result["dadata"] = run_sample(result)
            output = write_result(result, Path(tmp) / "output")
            saved = load_workbook(output)
            self.assertEqual(saved.sheetnames[0], "Школы и ученики")
            self.assertEqual(saved["Группы"]["B2"].value, "00012\n00013")
            self.assertEqual(saved["Группы"]["B2"].data_type, "s")
            saved.close()

    def test_personal_files_not_sent_to_dadata(self):
        result = process([record(1, "PRIVATE-1234567")], CONFIG)
        report = run_sample(result)
        self.assertNotIn("PRIVATE-1234567", json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    unittest.main()
