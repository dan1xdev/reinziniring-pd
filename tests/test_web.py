"""Проверки загрузки и границ локального HTTP-интерфейса."""
import base64
from http.client import HTTPConnection
from http.server import ThreadingHTTPServer
import io
import json
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from ou_processor.web import Application, Handler, JobCancelled, UserError, unpack_uploads


def upload(name='test.xlsx'):
    workbook = Workbook()
    workbook.active.append(['Произвольная колонка'])
    workbook.active.append(['Строка'])
    stream = io.BytesIO()
    workbook.save(stream)
    return {'name':name, 'data':base64.b64encode(stream.getvalue()).decode()}


def school_upload():
    workbook = Workbook()
    workbook.active.append(['Образовательное учреждение', 'Вид образования', 'ФИО', 'Личное дело.Номер'])
    workbook.active.append(['ГБОУ г. Москвы «Школа № 1»', 'Среднее общее', 'Иванов Иван', '0001'])
    stream = io.BytesIO()
    workbook.save(stream)
    return {'name':'schools.xlsx', 'data':base64.b64encode(stream.getvalue()).decode()}


class UploadTests(unittest.TestCase):
    def test_without_registry_processing_finishes_and_exports_review(self):
        class EmptyClient:
            def __init__(self, token):
                self.token = token
            def suggest(self, query):
                return []
        with tempfile.TemporaryDirectory() as root:
            app = Application(Path(root)/'runs', Path(root)/'registry', Path(root)/'cache.json',
                              token='test-key', client_factory=EmptyClient)
            job_id = app.start({'files':[school_upload()]})
            for _ in range(500):
                if app.snapshot(job_id)['status'] != 'running':
                    break
                time.sleep(.01)
            job = app.snapshot(job_id)
            self.assertEqual(job['status'], 'done', job.get('error'))
            self.assertEqual(job['summary']['review'], 1)
            self.assertTrue(any('Резервный реестр недоступен' in warning for warning in job['warnings']))
            book = load_workbook(Path(root)/'runs'/job_id/'result'/'schools.xlsx', read_only=True)
            self.assertEqual(book['Проверить']['E2'].value, 'Иванов Иван')
            self.assertEqual(book['Проверить']['F2'].value, '0001')
            book.close()

    def test_token_is_saved_and_reloaded_without_exposing_it(self):
        class Store:
            value = ''
            def load(self): return self.value
            def save(self, token): self.value = token
            def delete(self): self.value = ''
        with tempfile.TemporaryDirectory() as root:
            store = Store()
            app = Application(root, root, Path(root)/'cache.json', token_store=store)
            with patch.object(app, 'run'):
                app.start({'files':[upload()], 'token':'private-test-key', 'save_token':True})
            self.assertEqual(store.value, 'private-test-key')
            self.assertNotIn('private-test-key', json.dumps(app.settings()))
            reopened = Application(root, root, Path(root)/'cache.json', token_store=store)
            self.assertTrue(reopened.settings()['token_configured'])
            self.assertTrue(reopened.settings()['token_saved'])
            reopened.forget_token()
            self.assertEqual(store.value, '')

    def test_invalid_extension(self):
        with self.assertRaises(UserError): unpack_uploads([upload('test.xls')])

    def test_not_an_excel_archive(self):
        with self.assertRaises(UserError): unpack_uploads([{'name':'test.xlsx','data':'aGVsbG8='}])

    def test_duplicate_content_is_loaded_once(self):
        file = upload()
        self.assertEqual(len(unpack_uploads([file,{**file,'name':'copy.xlsx'}])),1)

    def test_filename_cannot_escape_upload_directory(self):
        self.assertEqual(unpack_uploads([upload('../../test.xlsx')])[0][0],'test.xlsx')

    def test_missing_key_is_reported(self):
        with tempfile.TemporaryDirectory() as root:
            app=Application(root,root,Path(root)/'cache.json')
            with self.assertRaisesRegex(UserError,'API-ключ'):
                app.start({'files':[upload()]})

    def test_job_failure_releases_worker(self):
        with tempfile.TemporaryDirectory() as root:
            app=Application(root,root,Path(root)/'cache.json')
            app.jobs['example']={'status':'running'}
            app.active='example'
            app.run('example',unpack_uploads([upload()]),'private-key',False)
            self.assertIsNone(app.active)
            self.assertEqual(app.snapshot('example')['status'],'error')
            self.assertNotIn('private-key',json.dumps(app.snapshot('example')))

    def test_second_job_is_rejected(self):
        with tempfile.TemporaryDirectory() as root:
            app=Application(root,root,Path(root)/'cache.json',token='key')
            app.active='existing'
            with patch.object(app,'settings',return_value={'registry_ready':True}):
                with self.assertRaisesRegex(UserError,'Уже идёт'):
                    app.start({'files':[upload()]})

    def test_cancel_stops_worker_and_removes_uploaded_file(self):
        entered, release = threading.Event(), threading.Event()
        def paused_read(paths, should_cancel):
            entered.set()
            release.wait(5)
            should_cancel()
        with tempfile.TemporaryDirectory() as root:
            app=Application(Path(root)/'runs',root,Path(root)/'cache.json',token='key')
            with patch.object(app,'settings',return_value={'registry_ready':True}), patch(
                'ou_processor.web.read_records',side_effect=paused_read):
                job_id=app.start({'files':[upload()]})
                self.assertTrue(entered.wait(5))
                self.assertEqual(app.cancel(job_id)['status'],'cancelling')
                with self.assertRaises(JobCancelled):
                    app.finish(job_id,summary={})
                release.set()
                for _ in range(200):
                    if app.snapshot(job_id)['status']=='cancelled':
                        break
                    time.sleep(.01)
                self.assertEqual(app.snapshot(job_id)['status'],'cancelled')
                self.assertIsNone(app.active)
                self.assertFalse((Path(root)/'runs'/job_id).exists())


class HTTPTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory()
        cls.server=ThreadingHTTPServer(('127.0.0.1',0),Handler)
        cls.server.app=Application(cls.temp.name,cls.temp.name,Path(cls.temp.name)/'cache.json',token='private-key')
        cls.thread=threading.Thread(target=cls.server.serve_forever,daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown(); cls.server.server_close(); cls.thread.join(); cls.temp.cleanup()

    def call(self,path,method='GET',body=None,headers=None):
        connection=HTTPConnection('127.0.0.1',self.server.server_port,timeout=5)
        connection.request(method,path,body,headers or {})
        response=connection.getresponse()
        status,data=response.status,response.read()
        connection.close()
        return status,data

    def test_config_does_not_expose_secret(self):
        status,data=self.call('/api/config')
        self.assertEqual(status,200)
        self.assertNotIn(b'private-key',data)
        self.assertTrue(json.loads(data)['token_configured'])

    def test_host_header_blocks_other_domains(self):
        status,_=self.call('/api/config',headers={'Host':'attacker.example'})
        self.assertEqual(status,403)

    def test_post_requires_csrf(self):
        status,_=self.call('/api/jobs','POST','{}')
        self.assertEqual(status,403)

    def test_cross_origin_is_rejected(self):
        status,_=self.call('/api/jobs','POST','{}',{'Origin':'https://other.example','X-CSRF-Token':self.server.app.csrf})
        self.assertEqual(status,403)

    def test_malformed_upload_is_explained(self):
        status,data=self.call('/api/jobs','POST','{"files":[]}',{'X-CSRF-Token':self.server.app.csrf})
        self.assertEqual(status,400)
        self.assertIn('Выберите',json.loads(data)['error'])

    def test_unknown_download_cannot_read_files(self):
        status,_=self.call('/api/jobs/../../README.txt/download')
        self.assertEqual(status,404)

    def test_cancel_endpoint_requires_csrf_and_stops_running_job(self):
        job_id='a'*32
        app=self.server.app
        app.jobs[job_id]={'id':job_id,'status':'running','phase':'Ищем организации','progress':30}
        app.cancel_flags[job_id]=threading.Event()
        try:
            path=f'/api/jobs/{job_id}/cancel'
            self.assertEqual(self.call(path,'POST')[0],403)
            status,data=self.call(path,'POST',headers={'X-CSRF-Token':app.csrf})
            self.assertEqual(status,202)
            self.assertEqual(json.loads(data)['status'],'cancelling')
            self.assertTrue(app.cancel_flags[job_id].is_set())
        finally:
            app.jobs.pop(job_id,None)
            app.cancel_flags.pop(job_id,None)

    def test_registry_retry_endpoint_starts_worker_without_blocking_request(self):
        app = self.server.app
        with patch.object(app, 'start_registry_update', return_value={
            'phase':'checking', 'attempt':1, 'max_attempts':3, 'message':'Проверяем',
            'retry_available':False}) as start:
            status, data = self.call('/api/registry/retry', 'POST', headers={'X-CSRF-Token':app.csrf})
        self.assertEqual(status, 202)
        self.assertEqual(json.loads(data)['attempt'], 1)
        start.assert_called_once_with()
