"""Регрессионные проверки правил и транспорта без внешних запросов."""

import copy
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock
from urllib.error import HTTPError, URLError

from openpyxl import Workbook, load_workbook

from ou_processor.dadata import DaDataClient, DaDataError, run_sample
from ou_processor.export import write_result
from ou_processor.input import read_records
from ou_processor.normalization import expanded_key, normalize, numbers
from ou_processor.pipeline import process


CONFIG = json.loads((Path(__file__).resolve().parents[1] / "config.json").read_text(encoding="utf-8"))


def row(number=1, **fields):
    return {"record_id": f"test:{number}", "source_file": "test.xlsx", "source_sheet": "Данные",
            "source_row": number + 1, "name": "ГБОУ города Москвы «Школа № 1»",
            "education": "Среднее общее", "formula_fields": [], **fields}


def candidate(number="1", status="ACTIVE"):
    return {"value": f"ГБОУ Школа № {number}", "data": {
        "inn": "0000000000", "kpp": "000000000", "ogrn": "0000000000000",
        "name": {"full_with_opf": f"ГБОУ Школа № {number}"}, "state": {"status": status},
        "address": {"value": "Тестовый адрес", "data": {"postal_code": "000001"}}}}


class NormalizationTests(unittest.TestCase):
    def test_case_quotes_spaces_and_number(self):
        self.assertEqual(normalize('  ШКОЛА «№12»  '), normalize('школа "№ 12"'))
        self.assertEqual(numbers("Школа №12А"), ("12а",))

    def test_letters_not_silently_replaced(self):
        self.assertNotEqual(normalize("Орёл"), normalize("Орел"))

    def test_expansion_only_in_search_variant(self):
        self.assertEqual(normalize("МБОУ СОШ"), "мбоу сош")
        self.assertIn("средняя общеобразовательная школа", expanded_key("МБОУ СОШ"))

    def test_original_name_preserved(self):
        name = '  ГБОУ   «Школа №1»  '
        self.assertEqual(process([row(name=name)], CONFIG)["records"][0]["name"], name)


class PipelineTests(unittest.TestCase):
    def result(self, *rows):
        return process(list(rows), copy.deepcopy(CONFIG))

    def test_higher_education_excluded(self):
        record = self.result(row(education="Высшее"))["records"][0]
        self.assertEqual(record["selection"], "ИСКЛЮЧЕНО")

    def test_college_at_university_remains(self):
        record = self.result(row(name="Университет — колледж", education="СПО"))["records"][0]
        self.assertEqual(record["selection"], "ОТОБРАНО")
        self.assertIn("ВОЗМОЖНОЕ_ПОДРАЗДЕЛЕНИЕ_ВУЗА", record["flags"])

    def test_document_fallback_not_training_level(self):
        record = self.result(row(education="", document="Аттестат о среднем общем образовании",
                                 training="Бакалавриат"))["records"][0]
        self.assertEqual(record["education_kind"], "Среднее общее")
        self.assertEqual(record["selection"], "ОТОБРАНО")

    def test_unknown_kind_held_for_review(self):
        record = self.result(row(education="Неизвестный диплом"))["records"][0]
        self.assertEqual(record["selection"], "ПРОВЕРИТЬ_ОТБОР")
        self.assertEqual(record["group_id"], "")

    def test_conflicting_education_held(self):
        record = self.result(row(document="Диплом о среднем профессиональном образовании"))["records"][0]
        self.assertEqual(record["selection"], "ПРОВЕРИТЬ_ОТБОР")

    def test_empty_name_held(self):
        self.assertEqual(self.result(row(name=""))["records"][0]["selection"], "ПРОВЕРИТЬ_ОТБОР")

    def test_same_name_case_groups(self):
        result = self.result(row(1), row(2, name=row()["name"].upper()))
        self.assertEqual(len(result["groups"]), 1)
        self.assertEqual(result["groups"][0]["row_count"], 2)

    def test_different_regions_separate(self):
        result = self.result(row(1, name="Школа № 1", region_raw="Курская обл."),
                             row(2, name="Школа № 1", region_raw="Ростовская обл."))
        self.assertEqual(len(result["groups"]), 2)
        self.assertEqual(result["possible_matches"], [])

    def test_missing_city_does_not_merge(self):
        result = self.result(row(1, name="Школа № 1", region_raw="Курская обл."),
                             row(2, name="Школа № 1", region_raw="Курская область"))
        self.assertEqual(len(result["groups"]), 2)

    def test_conflict_never_silently_prioritized(self):
        record = self.result(row(region_raw="Курская обл."))["records"][0]
        self.assertEqual(record["region"], "")
        self.assertIn("КОНФЛИКТ_РЕГИОНА", record["flags"])

    def test_region_hypothesis_not_fact(self):
        result = self.result(row(1, name="Школа № 1", region_raw="Курская область"),
                             row(2, name="Школа № 1"))
        record = result["records"][1]
        self.assertEqual(record["region"], "")
        self.assertEqual(record["region_hypothesis"], "Курская область")
        self.assertEqual(len(result["groups"]), 2)

    def test_no_hypothesis_when_regions_disagree(self):
        result = self.result(row(1, name="Школа № 1", region_raw="Курская область"),
                             row(2, name="Школа № 1", region_raw="Ростовская область"),
                             row(3, name="Школа № 1"))
        self.assertEqual(result["records"][2]["region_hypothesis"], "")

    def test_truncation_boundary(self):
        result = self.result(*[row(i, name="а" * length) for i, length in enumerate([149, 150, 151])])
        self.assertEqual(["ПОДОЗРЕНИЕ_НА_ОБРЕЗКУ" in r["flags"] for r in result["records"]], [False, True, False])

    def test_old_year_not_excluded(self):
        result = self.result(row(1, year="2021"), row(2, year="2022"))
        self.assertIn("СТАРЫЙ_ВЫПУСК", result["records"][0]["flags"])
        self.assertNotIn("СТАРЫЙ_ВЫПУСК", result["records"][1]["flags"])
        self.assertEqual(result["summary"]["selected_rows"], 2)

    def test_number_mismatch_not_suggested(self):
        result = self.result(row(1), row(2, name="ГБОУ города Москвы Школа №2"))
        self.assertEqual(result["possible_matches"], [])

    def test_abbreviation_is_candidate_not_merge(self):
        result = self.result(row(1), row(2, name="Государственное бюджетное общеобразовательное учреждение города Москвы Школа №1"))
        self.assertEqual(len(result["groups"]), 2)
        self.assertEqual(len(result["possible_matches"]), 1)
        self.assertEqual(result["summary"]["confirmed_organizations"], 0)

    def test_school_without_legal_prefix_is_candidate(self):
        result = self.result(row(1, name="ГБОУ г. Москвы Школа на Яузе"),
                             row(2, name="Школа на Яузе, г. Москва"))
        self.assertEqual(len(result["possible_matches"]), 1)


class InputExportTests(unittest.TestCase):
    def test_read_report_header_ignore_analyst_region_keep_student_locally(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "input.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "100 записей"
            sheet.append(["Название отчета"])
            sheet.append([])
            sheet.append(["Название ОУ в исходном Excel", "Вид образования", "Регион для поиска", "ФИО"])
            sheet.append(["  Школа №1  ", "Среднее общее", "Москва", "Не передавать"])
            other = workbook.create_sheet("Разбор 100")
            other.append(["Образовательное учреждение", "Вид образования"])
            other.append(["Школа №2", "Среднее общее"])
            workbook.save(path)
            records, sources = read_records([path, path])
            self.assertEqual(len(records), 1)
            self.assertEqual(records[0]["name"], "  Школа №1  ")
            self.assertEqual(records[0]["student_name"], "Не передавать")
            result = process(records, CONFIG)
            self.assertEqual(result["records"][0]["region"], "")
            self.assertNotIn("Не передавать", json.dumps(run_sample(result), ensure_ascii=False))
            self.assertEqual(sources[0]["skipped_sheets"], ["Разбор 100"])

    def test_input_formula_held(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "input.xlsx"
            workbook = Workbook()
            workbook.active.append(["Образовательное учреждение", "Вид образования"])
            workbook.active.append(['="Школа №1"', "Среднее общее"])
            workbook.save(path)
            records, _ = read_records([path])
            self.assertEqual(process(records, CONFIG)["records"][0]["selection"], "ПРОВЕРИТЬ_ОТБОР")

    def test_export_literals_and_leading_zero(self):
        result = process([row(name='=HYPERLINK("https://example.com")')], CONFIG)
        result["sources"] = []
        result["dadata"] = {"mode": "offline", "actual_requests": 0, "requests": [], "candidates": [
            {"candidate_inn": "0123456789", "postal_code": "000001"}]}
        with tempfile.TemporaryDirectory() as tmp:
            path = write_result(result, tmp)
            workbook = load_workbook(path)
            self.assertEqual(workbook["Записи"]["E2"].data_type, "s")
            self.assertEqual(workbook["Кандидаты DaData"]["D2"].value, "0123456789")
            self.assertEqual(workbook["Кандидаты DaData"]["J2"].value, "000001")
            workbook.close()


class DaDataTests(unittest.TestCase):
    def test_request_contract_and_no_personal_fields(self):
        opener = Mock(return_value=BytesIO(json.dumps({"suggestions": [candidate()]}).encode()))
        client = DaDataClient("test-token", opener=opener)
        client.suggest("Школа №1 Москва")
        request = opener.call_args.args[0]
        self.assertEqual(json.loads(request.data), {"query": "Школа №1 Москва", "count": 5, "type": "LEGAL"})
        self.assertEqual(request.get_header("Authorization"), "Token test-token")
        self.assertEqual(opener.call_args.kwargs["timeout"], 15)

    def test_success_stays_unconfirmed(self):
        result = process([row()], CONFIG)
        client = Mock()
        client.suggest.return_value = [candidate()]
        report = run_sample(result, 1, client, online=True, delay=0)
        self.assertEqual(report["actual_requests"], 1)
        self.assertEqual(report["requests"][0]["status"], "КАНДИДАТЫ_ДЛЯ_ПРОВЕРКИ")
        self.assertFalse(report["candidates"][0]["confirmed"])
        self.assertEqual(result["records"][0]["confirmed_inn"], "")

    def test_limit_and_no_key(self):
        result = process([row(i, name=f"ГБОУ Москвы Школа №{i}") for i in range(10)], CONFIG)
        report = run_sample(result, 3, online=True)
        self.assertEqual(len(report["requests"]), 3)
        self.assertEqual(report["actual_requests"], 0)
        self.assertTrue(all(r["status"] == "НЕ_ВЫПОЛНЕНО_НЕТ_КЛЮЧА" for r in report["requests"]))
        self.assertEqual(run_sample(result, 0)["requests"], [])

    def test_http_errors_redact_response_and_token(self):
        for code, expected in [(401, "ОШИБКА_АВТОРИЗАЦИИ"), (403, "ДОСТУП_ЗАПРЕЩЕН"),
                               (429, "ЛИМИТ_ЗАПРОСОВ"), (500, "HTTP_500")]:
            opener = Mock(side_effect=HTTPError("url", code, "secret-token", {}, None))
            with self.assertRaisesRegex(DaDataError, expected):
                DaDataClient("secret-token", opener=opener).suggest("Школа")

    def test_stop_after_denied(self):
        result = process([row(i, name=f"Школа №{i}") for i in range(3)], CONFIG)
        client = Mock()
        client.suggest.side_effect = DaDataError("ДОСТУП_ЗАПРЕЩЕН")
        report = run_sample(result, 3, client, online=True, delay=0)
        self.assertEqual(client.suggest.call_count, 1)
        self.assertEqual(report["actual_requests"], 1)
        self.assertEqual(report["requests"][-1]["status"], "ПРОПУЩЕНО_ПОСЛЕ_ОШИБКИ_ДОСТУПА")

    def test_network_error(self):
        client = DaDataClient("test", opener=Mock(side_effect=URLError("secret")))
        with self.assertRaisesRegex(DaDataError, "ОШИБКА_СЕТИ"):
            client.suggest("Школа")

    def test_bad_response(self):
        for response in [b"bad", b"{}", b"[]", b'{"suggestions":[{}]}']:
            client = DaDataClient("test", opener=Mock(return_value=BytesIO(response)))
            with self.assertRaisesRegex(DaDataError, "НЕВЕРНЫЙ_ФОРМАТ_ОТВЕТА"):
                client.suggest("Школа")

    def test_no_result_is_not_network_error(self):
        client = Mock()
        client.suggest.return_value = []
        report = run_sample(process([row()], CONFIG), 1, client, online=True, delay=0)
        self.assertEqual(report["requests"][0]["status"], "НЕ_НАЙДЕНО")

    def test_candidate_number_and_status_flags(self):
        client = Mock()
        client.suggest.return_value = [candidate("2", "LIQUIDATED")]
        report = run_sample(process([row()], CONFIG), 1, client, online=True, delay=0)
        self.assertIn("КОНФЛИКТ_НОМЕРА", report["candidates"][0]["flags"])
        self.assertIn("СТАТУС_НЕ_ПОДТВЕРЖДЕН_ДЕЙСТВУЮЩИМ", report["candidates"][0]["flags"])


if __name__ == "__main__":
    unittest.main()
